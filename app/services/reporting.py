from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List

import pandas as pd
from openpyxl.styles import Alignment, Font

from app.config import get_settings

settings = get_settings()


def _prepare_reports_dir() -> Path:
    reports_dir = settings.reports_dir
    reports_dir.mkdir(parents=True, exist_ok=True)
    return reports_dir


def generate_excel_report(
    metrics: Dict[str, Dict],
    top_sku: Iterable[Dict],
    date_from: datetime,
    date_to: datetime,
) -> Path:
    reports_dir = _prepare_reports_dir()
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = reports_dir / f"report-{date_from.date()}_{date_to.date()}_{timestamp}.xlsx"

    platforms = list(metrics.keys())
    rows: List[Dict] = []
    totals = {
        "orders_count": 0,
        "delivered": 0,
        "returns": 0,
        "revenue_delivered": 0.0,
        "cogs": 0.0,
        "commission": 0.0,
        "logistics": 0.0,
        "storage": 0.0,
        "ads": 0.0,
        "gross_profit": 0.0,
        "profit": 0.0,
    }
    for platform in platforms:
        data = metrics[platform]
        rows.append(
            {
                "Platform": platform,
                "Orders": data["orders_count"],
                "Delivered": data["delivered"],
                "Returns": data["returns"],
                "Revenue": data["revenue_delivered"],
                "COGS": data["cogs"],
                "Commission": data["commission"],
                "Logistics": data["logistics"],
                "Storage": data["storage"],
                "Ads": data["ads"],
                "Gross Profit": data["gross_profit"],
                "Profit": data["profit"],
                "ROMI": data["romi"],
            }
        )
        totals["orders_count"] += data["orders_count"]
        totals["delivered"] += data["delivered"]
        totals["returns"] += data["returns"]
        totals["revenue_delivered"] += data["revenue_delivered"]
        totals["cogs"] += data["cogs"]
        totals["commission"] += data["commission"]
        totals["logistics"] += data["logistics"]
        totals["storage"] += data["storage"]
        totals["ads"] += data["ads"]
        totals["gross_profit"] += data["gross_profit"]
        totals["profit"] += data["profit"]

    romi = totals["profit"] / totals["ads"] if totals["ads"] else None
    rows.append(
        {
            "Platform": "total",
            "Orders": totals["orders_count"],
            "Delivered": totals["delivered"],
            "Returns": totals["returns"],
            "Revenue": totals["revenue_delivered"],
            "COGS": totals["cogs"],
            "Commission": totals["commission"],
            "Logistics": totals["logistics"],
            "Storage": totals["storage"],
            "Ads": totals["ads"],
            "Gross Profit": totals["gross_profit"],
            "Profit": totals["profit"],
            "ROMI": romi,
        }
    )
    metrics_df = pd.DataFrame(rows)

    top_sku_df = pd.DataFrame(list(top_sku))

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        metrics_df.to_excel(writer, sheet_name="Metrics", index=False)
        top_sku_df.to_excel(writer, sheet_name="Top SKU", index=False)
    _style_excel(filename)
    return filename


def _style_excel(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        for column_cells in sheet.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = length + 2
        for row in sheet.iter_rows(min_row=1, max_row=1):
            for cell in row:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
    workbook.save(path)
